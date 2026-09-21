"""Check a final Windows distribution through its real embedded worker."""
from __future__ import annotations

import argparse
import json
import os
import subprocess
import tempfile
import textwrap
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path
from threading import Thread


def check_shell(worker):
    """Use the packaged PTY extension, then retain an ended session's output."""
    worker("""
        import os, tempfile, time
        from pathlib import Path
        from pycat.core.tools.process import BackgroundProcessManager, CommandExecutionRequest
        from pycat.models.contracts.config import ShellConfig
        manager = BackgroundProcessManager()
        with tempfile.TemporaryDirectory(prefix='shell 输入 ') as folder:
            root = Path(folder)
            (root / 'child').mkdir()
            config = ShellConfig(backend='cmd' if os.name == 'nt' else 'sh')
            def start():
                return manager.start(CommandExecutionRequest(command='', cwd=root,
                    session_root=root, conversation_id='release', interactive=True,
                    background=True), shell_config=config).process_id
            def send(identity, text, controller='agent'):
                manager.write(identity, text + '\\r', conversation_id='release', controller=controller)
            def wait_for(identity, text):
                deadline = time.monotonic() + 15
                while time.monotonic() < deadline:
                    if text in manager.read_logs(identity):
                        return
                    time.sleep(.03)
                raise AssertionError(manager.read_logs(identity))
            try:
                identity = start()
                send(identity, 'set PYCAT_RELEASE_INPUT=continued' if os.name == 'nt'
                     else 'PYCAT_RELEASE_INPUT=continued')
                send(identity, 'echo shell-%PYCAT_RELEASE_INPUT%' if os.name == 'nt'
                     else 'echo shell-$PYCAT_RELEASE_INPUT')
                wait_for(identity, 'shell-continued')
                send(identity, 'cd child')
                send(identity, 'cd' if os.name == 'nt' else 'pwd')
                wait_for(identity, str(root / 'child'))
                manager.resize(identity, 110, 35, conversation_id='release')
                manager.set_controller(identity, 'user', conversation_id='release')
                try:
                    send(identity, 'wrong controller')
                except PermissionError:
                    pass
                else:
                    raise AssertionError('Agent input allowed after user takeover')
                send(identity, 'exit 0', controller='user')
                assert manager.wait(identity, 15).exit_code == 0
                assert 'shell-continued' in manager.read_logs(identity)
                second = start()
                manager.kill(second, conversation_id='release')
                assert not manager.status(second).running
                assert {p.process_id for p in manager.list(conversation_id='release', include_exited=True)} == {identity, second}
            finally:
                manager.kill_all()
    """)


def check_tui(worker):
    worker("""
        import asyncio, tempfile
        from pycat import PyCat
        from pycat.tui.app import WorkbenchApp
        from textual.widgets import TextArea
        from pygments.lexers import get_lexer_by_name
        from pygments.styles import get_style_by_name
        assert get_lexer_by_name('python') and get_style_by_name('monokai')
        async def check():
            with tempfile.TemporaryDirectory() as directory:
                async with PyCat(data_dir=directory) as app:
                    terminal = WorkbenchApp(services=app._services, work_dir=directory)
                    async with terminal.run_test(size=(80, 24)) as pilot:
                        await pilot.pause()
                        composer = terminal.query_one('#composer', TextArea)
                        composer.load_text('release draft 中文')
                        assert composer.text == 'release draft 中文'
                        assert terminal.query_one('#transcript').size.height > 0
        asyncio.run(check())
    """)


def check_gui_shell(worker):
    worker("""
        import tempfile, time
        from types import SimpleNamespace
        from PyQt6.QtWidgets import QApplication, QWidget
        from PyQt6.QtCore import QThreadPool
        from pycat.core.app.container import AppContainer
        from pycat.gui.presenters.shell_presenter import ShellPresenter
        from pycat.gui.runtime.message_runtime import MessageRuntime
        from pycat.models.conversation import Conversation
        app = QApplication([])
        def until(predicate):
            deadline = time.monotonic() + 20
            while not predicate() and time.monotonic() < deadline:
                app.processEvents()
                time.sleep(.01)
            app.processEvents()
            assert predicate(), 'Shell window did not reach the expected state'
        with tempfile.TemporaryDirectory() as directory:
            container = AppContainer(data_dir=directory, background_loop=True, background_curation=False)
            host = QWidget()
            host.services = container.services
            conversation = container.services.conv_service.create('Release Shell')
            conversation.mode, conversation.work_dir = 'agent', directory
            conversation.settings = {'tool_approval': 'allow'}
            container.services.conv_service.save(conversation)
            host.current_conversation = conversation
            host.conversation_presenter = SimpleNamespace(refresh_processes=lambda: None)
            host.message_runtime = MessageRuntime(container.services.run_service, parent=host)
            presenter = ShellPresenter(host)
            try:
                presenter.open()
                presenter.new_terminal()
                until(lambda: presenter.window.process_id and presenter.window.view._input_enabled)
                identity = presenter.window.process_id
                host.current_conversation = Conversation(id='other', title='Other')
                presenter.write(identity, 'echo PYCAT_RELEASE_SHELL\\r')
                until(lambda: 'PYCAT_RELEASE_SHELL' in '\\n'.join(presenter.window.view.screen.display))
                assert presenter.window.conversation_id == conversation.id
                presenter.window.close()
                assert container.services.tools.processes(conversation.id)[0].running
                host.current_conversation = conversation
                presenter.open(identity)
                assert presenter.window.isVisible()
                presenter.stop(identity)
                until(lambda: not presenter.window.stop_button.isEnabled())
                assert not container.services.tools.processes(conversation.id)
                assert presenter.window.tabs.count() == 1
                assert '已退出' in presenter.window.state_label.text()
                assert 'PYCAT_RELEASE_SHELL' in '\\n'.join(presenter.window.view.screen.display)
                assert not presenter.window.view._input_enabled
                assert not presenter.window.grab().isNull()
            finally:
                presenter.dispose()
                container.close()
                QThreadPool.globalInstance().waitForDone(5000)
                host.deleteLater()
                app.processEvents()
    """)


def check_web(worker):
    """Exercise Uvicorn's dynamic imports and serve the bundled browser assets."""
    worker("""
        import socket, tempfile, threading, time
        import httpx, uvicorn
        from pycat.web.server import create_app
        with tempfile.TemporaryDirectory() as directory, socket.socket() as listener:
            listener.bind(('127.0.0.1', 0))
            port = listener.getsockname()[1]
            app = create_app(data_dir=directory, token='release-probe')
            server = uvicorn.Server(uvicorn.Config(app, log_level='error', loop='asyncio', lifespan='on'))
            thread = threading.Thread(target=server.run, kwargs={'sockets': [listener]}, daemon=True)
            thread.start()
            try:
                deadline = time.monotonic() + 15
                while not server.started and thread.is_alive() and time.monotonic() < deadline:
                    time.sleep(.03)
                assert server.started, 'Web server failed to start'
                with httpx.Client(base_url=f'http://127.0.0.1:{port}', trust_env=False) as client:
                    page = client.get('/')
                    assert page.status_code == 200 and 'PyCat' in page.text
                    assert client.get('/api/bootstrap').status_code == 401
                    response = client.get('/api/bootstrap', headers={'Authorization': 'Bearer release-probe'})
                    assert response.status_code == 200 and response.json()['protocol'] == 1, response.text
                    assert client.get('/brand.svg').status_code == 200
                    from pathlib import Path
                    import pycat.web.server as web_server
                    assets = Path(web_server.__file__).resolve().parents[1] / 'assets' / 'web'
                    for asset in assets.rglob('*'):
                        if asset.is_file():
                            response = client.get('/assets/' + asset.relative_to(assets).as_posix())
                            assert response.status_code == 200 and response.content == asset.read_bytes(), str(asset)
            finally:
                server.should_exit = True
                thread.join(15)
                assert not thread.is_alive(), 'Web server did not close'
    """)


def check_file_search(worker):
    """The published distribution must search without a developer's PATH rg."""
    worker("""
        import asyncio, hashlib, json, os, tempfile
        from pathlib import Path
        from pycat.core.tools.system.file_search import find_ripgrep, RIPGREP_BUNDLE
        from pycat.core.tools.system.filesystem import GrepTool
        from pycat.core.tools.base import ToolContext
        os.environ['PATH'] = ''
        assert find_ripgrep() == str(RIPGREP_BUNDLE), 'Bundled ripgrep is missing'
        manifest = json.loads((RIPGREP_BUNDLE.parent / 'manifest.json').read_text())
        for name, digest in manifest['files'].items():
            assert hashlib.sha256((RIPGREP_BUNDLE.parent / name).read_bytes()).hexdigest() == digest, name
        with tempfile.TemporaryDirectory() as temp:
            (Path(temp) / 'file with spaces.txt').write_text('bundle probe 42', encoding='utf-8')
            result = asyncio.run(GrepTool().execute({'query': 'probe [0-9]+', 'regex': True}, ToolContext(work_dir=temp)))
            assert not result.is_error, result.to_string()
            data = json.loads(result.to_string())
            assert data['backend'] == 'rg' and data['matches'][0]['text'] == 'bundle probe 42', data
    """)


def check_ssh_helper(worker):
    """Remote source must survive freezing and be usable by a system Python."""
    worker("""
        from pathlib import Path
        from pycat.core.hosts import ssh
        source = Path(ssh.__file__).with_name('remote_helper.py').read_bytes()
        namespace = {'__name__': 'remote_probe'}
        exec(compile(source, '<remote-probe>', 'exec'), namespace)
        assert namespace['PROTOCOL'] == ssh.remote_helper.PROTOCOL
        assert namespace['Helper']().handle('hello', {})['protocol'] == ssh.remote_helper.PROTOCOL
    """)


def check_search(worker):
    """Exercise the packaged DDGS implementation, engine discovery, parser and HTTP client."""
    class SearchPage(BaseHTTPRequestHandler):
        def do_POST(self):
            self.rfile.read(int(self.headers.get("Content-Length", 0)))
            page = b'<div class="body"><h2>PyCat search probe</h2><a href="https://example.com/pycat-search">Packaged search works</a></div>'
            self.send_response(200)
            self.send_header("Content-Type", "text/html; charset=utf-8")
            self.send_header("Content-Length", str(len(page)))
            self.end_headers()
            self.wfile.write(page)

        def log_message(self, *_args):
            pass

    with ThreadingHTTPServer(("127.0.0.1", 0), SearchPage) as server:
        thread = Thread(target=server.serve_forever, daemon=True)
        thread.start()
        try:
            worker(f"""
                import asyncio, os
                for key in ('DDGS_PROXY', 'HTTP_PROXY', 'HTTPS_PROXY', 'ALL_PROXY', 'http_proxy', 'https_proxy', 'all_proxy'):
                    os.environ.pop(key, None)
                from ddgs import DDGS
                search = DDGS(timeout=5)
                from ddgs.engines import ENGINES
                assert ENGINES.get('text'), 'No packaged search engines'
                assert search._get_engines('text', 'auto'), 'No usable packaged search engines'
                engine = ENGINES['text']['duckduckgo']
                engine.search_url = 'http://127.0.0.1:{server.server_port}/'
                ENGINES['text'] = {{'duckduckgo': engine}}
                from pycat.core.app.services.search import SearchService
                from pycat.models.search_config import SearchConfig
                result = asyncio.run(SearchService(SearchConfig(enabled=True)).search('PyCat search probe'))
                assert 'https://example.com/pycat-search' in result, result
                assert 'PyCat search probe' in result, result
            """)
        finally:
            server.shutdown()
            thread.join()


def check_mcp(worker, workspace: Path):
    """Exercise protocol fallback and a real tool call without external services."""
    server_script = workspace / "mcp stdio fixture.py"
    server_script.write_text(textwrap.dedent("""
        import json, sys
        for line in sys.stdin:
            request = json.loads(line)
            if 'id' not in request:
                continue
            if request['method'] == 'initialize':
                result = {'protocolVersion': request['params']['protocolVersion'],
                    'capabilities': {'tools': {}}, 'serverInfo': {'name': 'release-probe', 'version': '1'}}
            elif request['method'] == 'tools/list':
                result = {'tools': [{'name': 'echo', 'description': 'Release probe',
                    'inputSchema': {'type': 'object', 'properties': {}}}]}
            elif request['method'] == 'tools/call':
                result = {'content': [{'type': 'text', 'text': 'release echo'}], 'isError': False}
            else:
                print(json.dumps({'jsonrpc': '2.0', 'id': request['id'],
                    'error': {'code': -32601, 'message': 'Method not found'}}), flush=True)
                continue
            print(json.dumps({'jsonrpc': '2.0', 'id': request['id'], 'result': result}), flush=True)
    """), encoding="utf-8")
    worker(f"""
        import asyncio, tempfile
        from pycat import PyCat
        from pycat.models.contracts.mcp import McpServerConfig
        from pycat.core.tools.system.python_exec import resolve_python_runner
        async def check():
            with tempfile.TemporaryDirectory() as data:
                async with PyCat(data_dir=data) as app:
                    runner = resolve_python_runner()
                    config = McpServerConfig(name='probe', command=runner[0],
                        args=[*runner[1:], {str(server_script)!r}])
                    result = await app.mcp.probe(config)
                    assert result['ok'] and result['tools'] == ['echo'], result
                    assert app.mcp.save([config]).ok
                    result = await app.tools.call('mcp__probe__echo', {{}}, approval_callback=lambda _: True)
                    assert not result.is_error and result.content == 'release echo', result
        asyncio.run(check())
    """)


def check_export(worker):
    worker("""
        import tempfile
        from pathlib import Path
        from docx import Document
        from pycat import export_document
        with tempfile.TemporaryDirectory() as directory:
            target = Path(directory) / 'report.docx'
            text = '# Export probe\\n\\n**Formatted** text\\n\\n|Name|Result|\\n|---|---|\\n|Word|pass|'
            export_document(text, target)
            document = Document(target)
            assert document.paragraphs[0].text == 'Export probe'
            assert document.tables[0].cell(1, 1).text == 'pass'
            target = export_document(text, Path(directory) / 'report.html')
            assert '<table>' in target.read_text(encoding='utf-8')
    """)


def check_ocr(worker):
    worker("""
        import numpy as np
        import pymupdf
        from pycat.core.content.ocr import OcrService
        from pycat.core.content.ppocrv6 import PpOcrV6Engine
        tool = OcrService()
        assert tool.status().available, tool.status()
        engine = PpOcrV6Engine(
            detection_model=tool.assets_dir / "PP-OCRv6_det_small.onnx",
            recognition_model=tool.assets_dir / "PP-OCRv6_rec_small.onnx")
        assert engine.recognize(np.full((64, 128, 3), 255, dtype=np.uint8)) == []
        with pymupdf.open() as document:
            page = document.new_page(width=400, height=160)
            page.insert_text((30, 80), 'PyCat OCR 12345', fontsize=28)
            pixmap = page.get_pixmap(matrix=pymupdf.Matrix(2, 2), alpha=False)
            pixels = np.frombuffer(pixmap.samples, dtype=np.uint8).reshape(pixmap.height, pixmap.width, 3)
            text = ' '.join(block.text for block in engine.recognize(pixels))
            assert '12345' in text and 'pycat' in text.lower(), text
    """)


def check_binary(directory: Path, frontend: str, ocr: bool) -> list[str]:
    checks = []
    executable = directory / ("pycat.exe" if frontend == "gui" else "pycat-cli.exe")
    with tempfile.TemporaryDirectory(prefix="pycat release 测试 ") as temporary:
        workspace = Path(temporary)
        env = {**os.environ, "HOME": temporary, "USERPROFILE": temporary,
               "QT_QPA_PLATFORM": "offscreen", "PYTHONUTF8": "1", "PYTHONIOENCODING": "utf-8"}
        for key in ("PYTHONPATH", "PYTHONHOME", "PYCAT_PYTHON", "PYTHON"):
            env.pop(key, None)
        # The distribution must find its own DLLs, not a developer environment.
        windows = Path(os.environ.get("SystemRoot", r"C:\Windows"))
        env["PATH"] = os.pathsep.join([str(windows / "System32"), str(windows)])
        for key in ("QT_PLUGIN_PATH", "QT_QPA_PLATFORM_PLUGIN_PATH", "QML2_IMPORT_PATH"):
            env.pop(key, None)

        def invoke(args, *, binary=executable, input="", expected=0):
            result = subprocess.run(
                [str(binary), *args], cwd=workspace, env=env, input=input,
                text=True, encoding="utf-8", errors="replace", capture_output=True, timeout=120,
                creationflags=getattr(subprocess, 'CREATE_NO_WINDOW', 0),
            )
            if result.returncode != expected:
                raise RuntimeError(f"{binary.name} {args[0]}: exit {result.returncode}\n"
                                   f"{result.stdout}\n{result.stderr}")
            return result

        def worker(code, **options):
            script = workspace / "worker quoted path 测试.py"
            script.write_text(textwrap.dedent(code), encoding="utf-8")
            return invoke(["--pycat-python-exec-worker", str(script)], **options)

        if frontend != "gui":
            assert invoke(["--version"]).stdout.startswith("PyCat ")
            assert json.loads(invoke(["mode", "list", "--output-format", "json"]).stdout)
            checks.append("CLI version, JSON and isolated profile")
        result = worker("""
            import json, sys
            print(json.dumps({"input": sys.stdin.read()}, ensure_ascii=True))
            print("diagnostic", file=sys.stderr)
            raise SystemExit(7)
        """, input="pipe input", expected=7)
        assert json.loads(result.stdout) == {"input": "pipe input"} and result.stderr.strip() == "diagnostic"
        checks.append("embedded worker, Unicode paths, stdin/stdout/stderr, exit code")
        worker("""
            from pathlib import Path
            import subprocess, tempfile
            from pycat.core.tools.system.python_exec import resolve_python_runner
            runner = resolve_python_runner()
            assert Path(runner[0]).is_file(), runner
            with tempfile.TemporaryDirectory() as folder:
                script = Path(folder) / 'nested.py'
                script.write_text("print('nested worker')\\nraise SystemExit(9)\\n", encoding='utf-8')
                result = subprocess.run([*runner, str(script)], input='', capture_output=True, text=True, timeout=20)
                assert result.returncode == 9 and result.stdout.strip() == 'nested worker', result
        """)
        checks.append("worker resolves its own executable before MCP/AnyIO imports and can spawn a nested worker")
        worker("""
            import json, os, tempfile
            from importlib.resources import files
            import psutil
            from pycat.core.skills.discovery import SkillsManager
            assert psutil.Process(os.getpid()).is_running()
            config = files("pycat").joinpath("assets", "extensions", "browser.json")
            assert isinstance(json.loads(config.read_text(encoding="utf-8")), dict)
            with tempfile.TemporaryDirectory() as data:
                skills = SkillsManager(work_dir=data, data_dir=data)
                for name in ("find-skills", "skill-creator"):
                    skill = skills.get(name)
                    assert skill and skill.source_scope == "bundled" and skill.read_only, name
                    assert skill.description and skill.content, name
        """)
        checks.append("Bundled skill discovery, browser config and native process inspection")
        worker("""
            import asyncio, io, json, tempfile
            from importlib.resources import files
            from pathlib import Path
            import httpx, pymupdf, ssl
            from PIL import Image
            from pycat import PyCat, Provider, ModelProfile, RunRequest
            from pycat.core.content.ocr import OcrService
            seed = files("pycat").joinpath("assets/default_models.json")
            assert json.loads(seed.read_text(encoding="utf-8"))["providers"]
            assert not files("pycat").joinpath("core/app/default_models.json").is_file()
            async def respond(request):
                payload = {"choices": [{"delta": {"content": "compiled answer"}, "finish_reason": "stop"}]}
                return httpx.Response(200, headers={"content-type": "text/event-stream"},
                    text="data: " + json.dumps(payload) + "\\n\\ndata: [DONE]\\n\\n")
            async def run():
                with tempfile.TemporaryDirectory() as data:
                    async with PyCat(data_dir=data, transport_factory=lambda: httpx.MockTransport(respond)) as app:
                        assert app.modes.list()
                        app.models.save([Provider(name="test", api_base="https://model.invalid/v1",
                            models=[ModelProfile(model_id="test")])])
                        first = await app.run(RunRequest(text="one"))
                        second = await app.run(RunRequest(text="two", conversation_id=first.conversation.id))
                        assert second.final_message.content == "compiled answer"
                        assert len(app.conversations.load(first.conversation.id).messages) == 4
            asyncio.run(run())
            document = pymupdf.open()
            document.new_page().insert_text((72, 72), "PDF probe")
            with pymupdf.open(stream=document.tobytes(), filetype="pdf") as reopened:
                assert "PDF probe" in reopened[0].get_text()
                png = reopened[0].get_pixmap().tobytes("png")
                with Image.open(io.BytesIO(png)) as rendered:
                    assert rendered.width > 100 and rendered.height > 100
            document.close()
            print("SDK and PDF passed")
        """)
        checks.append("SDK two-turn persistence, model transport, PDF and image runtime")
        check_ssh_helper(worker)
        checks.append("SSH helper source, protocol and stdlib bootstrap resource")
        check_export(worker)
        checks.append("DOCX templates, formatted text/table round trip and HTML export")
        check_search(worker)
        checks.append("DDGS dynamic engines, native HTTP client, HTML parsing and search service (local fixture)")
        check_mcp(worker, workspace)
        checks.append("MCP 2 SDK, legacy negotiation, discovery and tool call through a real stdio child")
        check_file_search(worker)
        checks.append("Pinned ripgrep, licenses, integrity and real regex search without PATH")
        check_shell(worker)
        checks.append("Native PTY, later input, working directory, resize, control transfer, exit/stop and retained output")
        if frontend != "gui":
            check_tui(worker)
            checks.append("TUI startup, styles, dynamic widgets, syntax lexers, Unicode input and shutdown")
            check_web(worker)
            checks.append("Local Uvicorn server, Web authentication, bootstrap, bundled assets and shutdown")
        if ocr:
            check_ocr(worker)
            assert not list(directory.rglob("opencv_videoio_ffmpeg*.dll"))
            assert not list(directory.rglob("onnxruntime.dll"))
            checks.append("OCR DLLs, both model sessions, blank/text inference; unused video and standalone C API DLLs absent")
        else:
            worker("""
                from importlib.util import find_spec
                from pycat.core.content.ocr import OcrService
                assert all(find_spec(n) is None for n in ("numpy", "onnxruntime", "cv2", "pyclipper"))
                assert not OcrService().status().available
            """)
            checks.append("OCR omission is explicit and reported unavailable")
        if frontend != "cli":
            check_gui_shell(lambda code: worker(code, binary=directory / "pycat.exe"))
            checks.append("GUI Shell creation, native output, pinned session, hide/reopen and retained stopped tab")
            worker("""
                import threading
                from PyQt6.QtWidgets import QApplication
                from PyQt6.QtGui import QImageReader
                from PyQt6.QtCore import QThreadPool, QTimer
                from pycat.gui.application import _install_qtbase_translation
                from pycat.gui.main_window import MainWindow
                from pycat.gui.runtime.background_job import BackgroundJob
                from pycat.gui.widgets.terminal_view import TerminalView
                from pycat.gui.runtime.content_navigation import ContentOpenUseCase
                import base64
                app = QApplication([])
                formats = {bytes(fmt).decode() for fmt in QImageReader.supportedImageFormats()}
                assert {'png', 'jpeg', 'webp', 'gif', 'svg'} <= formats, formats
                svg = b'<svg xmlns="http://www.w3.org/2000/svg" width="40" height="20"><rect width="40" height="20" fill="#16794c"/></svg>'
                image = ContentOpenUseCase.read_image('data:image/svg+xml;base64,' + base64.b64encode(svg).decode())
                assert image.width() == 40 and image.pixelColor(10, 10).name() == '#16794c'
                assert _install_qtbase_translation(app)
                window = MainWindow()
                window.show()
                terminal = TerminalView(window)
                terminal.resize(640, 320)
                terminal.feed('\\x1b[32mPyCat terminal 中文\\x1b[0m')
                assert 'PyCat terminal 中文' in '\\n'.join(terminal.screen.display)
                assert not terminal.grab().isNull()
                owner_thread = threading.get_ident()
                delivered = []
                def finished(result, error):
                    delivered.append((result, error, threading.get_ident()))
                    window.close()
                job = BackgroundJob(threading.get_ident)
                job.signals.finished.connect(finished)
                QThreadPool.globalInstance().start(job)
                QTimer.singleShot(10000, app.quit)
                app.exec()
                QThreadPool.globalInstance().waitForDone(5000)
                assert len(delivered) == 1 and delivered[0][1] is None, delivered
                assert delivered[0][0] != owner_thread and delivered[0][2] == owner_thread, delivered
                assert not window.isVisible()
            """, binary=directory / "pycat.exe")
            checks.append("GUI entry, MainWindow, Qt plugins, translation, background thread delivery and close")
            assert not list(directory.rglob("qpdf.dll")) and not list(directory.rglob("Qt6Pdf.dll"))
            assert not list(directory.rglob("*WebEngine*"))
        if frontend == "all":
            result = worker("""
                import ctypes, json, sys
                assert ctypes.windll.kernel32.GetConsoleWindow() == 0
                print(json.dumps({"input": sys.stdin.read()}))
                raise SystemExit(7)
            """, binary=directory / "pycat.exe", input="GUI pipe", expected=7)
            assert json.loads(result.stdout) == {"input": "GUI pipe"}
            checks.append("direct GUI entry hides console and preserves pipes/exit")
        if frontend == "cli":
            worker("from importlib.util import find_spec; assert find_spec('PyQt6') is None")
            result = invoke(["--gui"], expected=2)
            assert "does not include the desktop frontend" in result.stderr
            checks.append("headless distribution excludes Qt and explains its missing GUI")
    return checks


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("directory", type=Path)
    parser.add_argument("--frontend", choices=("all", "gui", "cli"), default="all")
    parser.add_argument("--without-ocr", action="store_true")
    options = parser.parse_args()
    print(json.dumps(check_binary(options.directory.resolve(), options.frontend, not options.without_ocr),
                     ensure_ascii=False, indent=2))
