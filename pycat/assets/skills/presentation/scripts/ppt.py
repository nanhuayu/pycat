"""Local presentation CLI for the existing Python/Shell execution tools."""
from __future__ import annotations

import argparse
import hashlib
import importlib.metadata
import importlib.util
import json
import os
import platform
import sys
import tempfile
import zipfile
from pathlib import Path

sys.dont_write_bytecode = True
sys.path.insert(0, str(Path(__file__).resolve().parent))

MAX_SOURCE_BYTES = 256 * 1024
MAX_DECK_BYTES = 20 * 1024 * 1024
MAX_SLIDES = 128
MAX_NODES = 10_000


class ScriptError(ValueError):
    def __init__(self, code: str, message: str, **details):
        super().__init__(message)
        self.code = code
        self.details = details


class Parser(argparse.ArgumentParser):
    def error(self, message):
        raise ScriptError("invalid_arguments", message)


def doctor() -> dict:
    dependencies = {}
    for distribution, module in (("python-pptx", "pptx"), ("Pillow", "PIL"), ("qrcode", "qrcode"),
                                 ("lxml", "lxml"), ("XlsxWriter", "xlsxwriter")):
        available = importlib.util.find_spec(module) is not None
        try:
            version = importlib.metadata.version(distribution)
        except importlib.metadata.PackageNotFoundError:
            version = "bundled" if available else "missing"
        dependencies[distribution] = {"available": available, "version": version}
    return {"ok": all(item["available"] for item in dependencies.values()),
            "dependencies": dependencies,
            "runtime": {"host": "local", "platform": platform.system(), "python": platform.python_version()},
            "commands": ["doctor", "validate", "build", "inspect"]}


def workspace_path(value: str) -> Path:
    root = Path.cwd().resolve()
    candidate = (root / value).resolve()
    try:
        candidate.relative_to(root)
    except ValueError as exc:
        raise ScriptError("invalid_path", "Path must stay inside the current workspace.") from exc
    return candidate


def digest(path: Path) -> str | None:
    if not path.is_file():
        return None
    with path.open("rb") as stream:
        return hashlib.file_digest(stream, "sha256").hexdigest()


def read_source(path: Path, project: Path) -> str:
    try:
        path.resolve().relative_to(project)
    except ValueError as exc:
        raise ScriptError("invalid_path", "Slide source escapes the project directory.") from exc
    if path.stat().st_size > MAX_SOURCE_BYTES:
        raise ScriptError("resource_limit", "Slide source exceeds 256 KiB.", source=path.name)
    with path.open("rb") as stream:
        content = stream.read(MAX_SOURCE_BYTES + 1)
    if len(content) > MAX_SOURCE_BYTES:
        raise ScriptError("resource_limit", "Slide source exceeds 256 KiB.", source=path.name)
    try:
        return content.decode("utf-8-sig")
    except UnicodeError as exc:
        raise ScriptError("invalid_source", "Slide sources must be UTF-8.", source=path.name) from exc


def load_project(value: str):
    from pyppt.dsl import Element, parse_document
    from pyppt.lint import Linter

    project = workspace_path(value)
    if not project.is_dir():
        raise ScriptError("invalid_project", "Project directory not found.")
    folder = project / "slides"
    sources = sorted((folder if folder.is_dir() else project).glob("*.slide"))
    if not sources or len(sources) > MAX_SLIDES:
        raise ScriptError("invalid_project", "Expected 1–128 .slide files.")
    slides, diagnostics = [], []
    for source in sources:
        text = read_source(source, project)
        try:
            parsed = parse_document(text)
            pending = [(item, 0) for item in parsed.slides]
            count = 0
            while pending:
                element, depth = pending.pop()
                count += 1
                if count > MAX_NODES or depth > 64:
                    raise ScriptError("resource_limit", "Slide exceeds the node/depth budget.", source=source.name)
                pending.extend((child, depth + 1) for child in element.children if isinstance(child, Element))
            result = Linter(project).lint_text(text, source.name)
        except RecursionError as exc:
            raise ScriptError("resource_limit", "Slide exceeds the parser depth budget.", source=source.name) from exc
        except ScriptError:
            raise
        except (TypeError, ValueError, OverflowError) as exc:
            raise ScriptError("validation_failed", str(exc), source=source.name) from exc
        diagnostics.append({"source": source.name, **result.to_dict()})
        if result.ok:
            slides.append(parsed.slides[0])
    failed = [item for item in diagnostics if not item["ok"]]
    if failed:
        raise ScriptError("validation_failed", "Slide validation failed; no output committed.", diagnostics=failed[:20])
    return project, slides, diagnostics


def validate(value: str) -> dict:
    _, slides, diagnostics = load_project(value)
    return {"ok": True, "slides": len(slides), "diagnostics": diagnostics[:128],
            "layout_measurement": "estimated; final fonts and rendering require visual review"}


def check_package(path: Path) -> None:
    if not path.is_file() or path.suffix.lower() != ".pptx":
        raise ScriptError("invalid_source", "Expected an existing .pptx file.")
    if path.stat().st_size > MAX_DECK_BYTES:
        raise ScriptError("resource_limit", "PPTX exceeds the 20 MiB budget.")
    try:
        with zipfile.ZipFile(path) as package:
            parts = package.infolist()
            if len(parts) > 4096 or sum(item.file_size for item in parts) > 128 * 1024 * 1024:
                raise ScriptError("resource_limit", "PPTX package exceeds the entry/expanded-size budget.")
            if any(item.file_size > 8 * 1024 * 1024 for item in parts if item.filename.endswith((".xml", ".rels"))):
                raise ScriptError("resource_limit", "PPTX XML part exceeds 8 MiB.")
    except zipfile.BadZipFile as exc:
        raise ScriptError("invalid_source", "Invalid PPTX package.") from exc


def verify_output(path: Path, expected_pages: int) -> None:
    from pptx import Presentation

    check_package(path)
    prs = Presentation(path)
    if len(prs.slides) != expected_pages:
        raise ScriptError("invalid_output", "Generated page count does not match the sources.")
    for slide in prs.slides:
        for shape in slide.shapes:
            if shape.shape_type == 13:
                _ = shape.image.blob
            if shape.has_chart:
                _ = [series.values for series in shape.chart.series]
            if shape.has_text_frame:
                for paragraph in shape.text_frame.paragraphs:
                    for run in paragraph.runs:
                        _ = run.hyperlink.address


def build(project_value: str, output_value: str, expected: str | None) -> dict:
    from pyppt.config import Settings
    from pyppt.render import Renderer

    project, slides, diagnostics = load_project(project_value)
    output = workspace_path(output_value)
    if output.suffix.lower() != ".pptx":
        raise ScriptError("invalid_path", "Output must use the .pptx extension.")
    output.parent.mkdir(parents=True, exist_ok=True)
    lock = output.with_name(output.name + ".lock")
    try:
        handle = os.open(lock, os.O_CREAT | os.O_EXCL | os.O_WRONLY)
    except FileExistsError as exc:
        raise ScriptError("output_busy", "Another build is using this output.") from exc
    os.close(handle)
    temporary = None
    try:
        before = digest(output)
        if before != expected:
            raise ScriptError("output_conflict", "Output version changed or overwrite was not explicitly requested.", current_digest=before)
        with tempfile.NamedTemporaryFile(dir=output.parent, prefix="." + output.stem + "-",
                                         suffix=".tmp.pptx", delete=False) as stream:
            temporary = Path(stream.name)
        Renderer(Settings(workspace=project), project).render_deck(slides, temporary)
        verify_output(temporary, len(slides))
        if digest(output) != before:
            raise ScriptError("output_conflict", "Output changed during rendering; existing file preserved.")
        os.replace(temporary, output)
        temporary = None
        return {"ok": True, "path": str(output), "digest": digest(output), "slides": len(slides),
                "size_bytes": output.stat().st_size,
                "warnings": [{"source": item["source"], "warnings": item["warnings"]}
                             for item in diagnostics if item["warnings"]],
                "preview": "PPTX structure verified; no raster preview generated",
                "delivery": "Use file__deliver for the committed workspace file."}
    finally:
        if temporary is not None:
            temporary.unlink(missing_ok=True)
        lock.unlink(missing_ok=True)


def inspect(value: str, start_page: int, page_count: int) -> dict:
    from pptx import Presentation

    path = workspace_path(value)
    check_package(path)
    prs = Presentation(path)
    total = len(prs.slides)
    if not 1 <= start_page <= total or not 1 <= page_count <= 10:
        raise ScriptError("invalid_arguments", "Use valid 1-based pages and page-count between 1 and 10.")
    end = min(total, start_page - 1 + page_count)
    pages, budget = [], 64_000
    for index in range(start_page - 1, end):
        slide = prs.slides[index]
        chunks, shapes = [], []
        for shape in slide.shapes:
            if shape.has_text_frame:
                chunks.append(shape.text)
            if shape.has_table:
                chunks.extend(" | ".join(cell.text for cell in row.cells) for row in shape.table.rows)
            if len(shapes) < 200:
                shapes.append({"id": shape.shape_id, "name": shape.name[:200], "type": str(shape.shape_type),
                               "x": shape.left, "y": shape.top, "width": shape.width, "height": shape.height})
        text = "\n".join(chunks)
        notes = slide.notes_slide.notes_text_frame.text if slide.has_notes_slide else ""
        allowance = min(8000, budget)
        shown_text = text[:allowance]
        shown_notes = notes[:min(2000, max(0, budget - len(shown_text)))]
        budget -= len(shown_text) + len(shown_notes)
        pages.append({"page": index + 1, "text": shown_text, "notes": shown_notes,
                      "text_truncated": len(shown_text) < len(text) or len(shown_notes) < len(notes),
                      "shapes": shapes, "shapes_truncated": len(slide.shapes) > len(shapes)})
    return {"ok": True, "path": str(path), "digest": digest(path), "total_pages": total,
            "pages": pages, "geometry_unit": "EMU", "next_page": end + 1 if end < total else None}


def parser() -> Parser:
    result = Parser(description="Local PPTX build and bounded inspection.")
    commands = result.add_subparsers(dest="command", required=True)
    commands.add_parser("doctor")
    check = commands.add_parser("validate")
    check.add_argument("project")
    create = commands.add_parser("build")
    create.add_argument("project")
    create.add_argument("--output", required=True)
    create.add_argument("--expected-digest")
    read = commands.add_parser("inspect")
    read.add_argument("path")
    read.add_argument("--start-page", type=int, default=1)
    read.add_argument("--page-count", type=int, default=5)
    return result


def main(argv=None) -> int:
    try:
        args = parser().parse_args(argv)
        dependencies = doctor()
        if args.command == "doctor":
            receipt = dependencies
        else:
            if not dependencies["ok"]:
                raise ScriptError("missing_dependencies", "Use a verified interpreter or a release containing these dependencies.", dependencies=dependencies["dependencies"])
            if args.command == "validate":
                receipt = validate(args.project)
            elif args.command == "build":
                receipt = build(args.project, args.output, args.expected_digest)
            else:
                receipt = inspect(args.path, args.start_page, args.page_count)
    except ScriptError as exc:
        receipt = {"ok": False, "error_code": exc.code, "message": str(exc), **exc.details}
    except ImportError as exc:
        receipt = {"ok": False, "error_code": "missing_dependencies", "message": str(exc)[:4000]}
    except (OSError, ValueError, KeyError, TypeError, OverflowError, SyntaxError, RecursionError) as exc:
        receipt = {"ok": False, "error_code": "operation_failed", "message": str(exc)[:4000]}
    print(json.dumps(receipt, ensure_ascii=False))
    return 0 if receipt["ok"] else 1


if __name__ == "__main__":
    raise SystemExit(main())
