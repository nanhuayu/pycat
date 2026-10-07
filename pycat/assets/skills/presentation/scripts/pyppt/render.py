"""渲染层：布局结果 → .pptx（python-pptx）。

覆盖 SlideDSL 组件：
Slide / Box / Text(span/br) / Image(Picture) / Table / Chart / FAIcon /
SVG / QRCode / CodeBlock，以及 background / border / borderRadius /
gradient / shadow / opacity / href / notes。
"""

from __future__ import annotations

import io
import logging
import re
from pathlib import Path
from typing import Any, Dict, List, Optional, Sequence, Tuple

import qrcode
from PIL import Image, ImageDraw
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Pt

from .config import CANVAS_H, CANVAS_W, MAX_ASSET_BYTES, Settings
from .dsl import Element
from .layout import Box, layout_slide, runs_of
from .style import Color, parse_border, parse_color, parse_gradient, parse_padding, resolve_length
from .units import px_to_emu

log = logging.getLogger("pyppt.render")

def _ms(name: str):
    """安全查 MSO_SHAPE 枚举（不同 python-pptx 版本枚举名有差异）。"""
    return getattr(MSO_SHAPE, name, None)


# Font Awesome → PowerPoint 自选图形 回退表（缺失的枚举自动降级为字形）
ICON_SHAPES: Dict[str, Any] = {
    "check": None,
    "check-circle": _ms("OVAL"),
    "check-square": _ms("ROUNDED_RECTANGLE"),
    "times": _ms("MATH_MULTIPLY"),
    "times-circle": _ms("OVAL"),
    "exclamation-circle": _ms("OVAL"),
    "info-circle": _ms("OVAL"),
    "arrow-right": _ms("RIGHT_ARROW"),
    "arrow-left": _ms("LEFT_ARROW"),
    "arrow-up": _ms("UP_ARROW"),
    "arrow-down": _ms("DOWN_ARROW"),
    "chevron-right": _ms("CHEVRON"),
    "plus": _ms("MATH_PLUS"),
    "minus": _ms("MATH_MINUS"),
    "star": _ms("STAR_5_POINT"),
    "heart": _ms("HEART"),
    "home": _ms("PENTAGON"),
    "user": _ms("OVAL"),
    "users": _ms("OVAL"),
    "file": _ms("FOLDED_CORNER"),
    "file-text": _ms("FOLDED_CORNER"),
    "folder": _ms("FOLDED_CORNER"),
    "copy": _ms("FLOWCHART_DOCUMENT"),
    "clipboard": _ms("FLOWCHART_DOCUMENT"),
    "briefcase": _ms("ROUNDED_RECTANGLE"),
    "chart-line": _ms("RECTANGLE"),
    "chart-bar": _ms("RECTANGLE"),
    "dollar": _ms("OVAL"),
    "calendar": _ms("ROUNDED_RECTANGLE"),
    "clock": _ms("OVAL"),
    "lock": _ms("ROUNDED_RECTANGLE"),
    "unlock": _ms("ROUNDED_RECTANGLE"),
    "shield": _ms("PENTAGON"),
    "key": _ms("OVAL"),
    "eye": _ms("OVAL"),
    "cog": _ms("GEAR_6"),
    "wrench": _ms("ROUNDED_RECTANGLE"),
    "edit": None,
    "pencil": None,
    "trash": _ms("ROUNDED_RECTANGLE"),
    "envelope": None,
    "phone": _ms("ROUNDED_RECTANGLE"),
    "comment": _ms("ROUNDED_RECTANGLE"),
    "share": _ms("CIRCULAR_ARROW"),
    "rocket": _ms("ISOSCELES_TRIANGLE"),
    "lightbulb": _ms("OVAL"),
    "search": _ms("OVAL"),
    "settings": _ms("GEAR_6"),
    "download": _ms("DOWN_ARROW"),
    "upload": _ms("UP_ARROW"),
    "link": _ms("ROUNDED_RECTANGLE"),
    "play": _ms("ISOSCELES_TRIANGLE"),
    "pause": _ms("RECTANGLE"),
    "stop": _ms("OCTAGON"),
    "flag": _ms("PENTAGON"),
    "bookmark": _ms("PENTAGON"),
    "tag": _ms("PENTAGON"),
    "globe": _ms("OVAL"),
    "database": _ms("CAN"),
    "server": _ms("CAN"),
    "cloud": _ms("CLOUD"),
    "code": _ms("ROUNDED_RECTANGLE"),
    "terminal": _ms("ROUNDED_RECTANGLE"),
    "bolt": _ms("LIGHTNING_BOLT"),
    "fire": _ms("OVAL"),
    "thumbs-up": _ms("ROUNDED_RECTANGLE"),
    "hand-point-right": _ms("RIGHT_ARROW"),
    "close": _ms("MATH_MULTIPLY"),
    "cross": _ms("CROSS"),
    "circle": _ms("OVAL"),
    "square": _ms("RECTANGLE"),
    "diamond": _ms("DIAMOND"),
}
ICON_SHAPES = {k: v for k, v in ICON_SHAPES.items() if v is not None}
ICON_GLYPHS: Dict[str, str] = {
    "heart": "\u2665",
    "dollar": "$",
    "info-circle": "i",
    "exclamation-circle": "!",
    "times": "\u00d7",
    "times-circle": "\u00d7",
    "plus": "+",
    "minus": "\u2212",
    "check": "\u2713",
    "check-circle": "\u2713",
    "check-square": "\u2713",
    "search": "\u2315",
    "star": "\u2605",
    "clock": "\u25f4",
    "bullet": "\u25cf",
}
CHART_TYPES = {
    "line": XL_CHART_TYPE.LINE_MARKERS,
    "bar": XL_CHART_TYPE.COLUMN_CLUSTERED,
    "column": XL_CHART_TYPE.COLUMN_CLUSTERED,
    "column-stacked": XL_CHART_TYPE.COLUMN_STACKED,
    "bar-horizontal": XL_CHART_TYPE.BAR_CLUSTERED,
    "pie": XL_CHART_TYPE.PIE,
    "doughnut": XL_CHART_TYPE.DOUGHNUT,
    "area": XL_CHART_TYPE.AREA,
    "scatter": XL_CHART_TYPE.XY_SCATTER,
}


def resolve_chart_type(chart_type: str, bar_direction: str = "column", grouping: str = "") -> Any:
    """把规范中的 chartType / barDirection / grouping 映射为 python-pptx 枚举。"""
    name = (chart_type or "").strip().lower()
    stacked = grouping in ("stacked", "stacked100", "stacked_100", "percentstacked")
    if name in ("barchart", "bar", "bar3dchart"):
        if bar_direction == "bar":
            return XL_CHART_TYPE.BAR_STACKED if stacked else XL_CHART_TYPE.BAR_CLUSTERED
        return XL_CHART_TYPE.COLUMN_STACKED if stacked else XL_CHART_TYPE.COLUMN_CLUSTERED
    if name in ("linechart", "line"):
        return XL_CHART_TYPE.LINE_MARKERS_STACKED if stacked else XL_CHART_TYPE.LINE_MARKERS
    if name in ("areachart", "area"):
        return XL_CHART_TYPE.AREA_STACKED if stacked else XL_CHART_TYPE.AREA
    if name in ("piechart", "pie"):
        return XL_CHART_TYPE.PIE
    if name in ("doughnutchart", "doughnut", "donut"):
        return XL_CHART_TYPE.DOUGHNUT
    if name in ("scatter", "scatterchart", "xyscatter"):
        return XL_CHART_TYPE.XY_SCATTER
    if name in ("radar", "radarchart"):
        return XL_CHART_TYPE.RADAR
    fallback = CHART_TYPES.get(name)
    return fallback if fallback is not None else XL_CHART_TYPE.COLUMN_CLUSTERED
ALIGN_MAP = {
    "left": PP_ALIGN.LEFT, "center": PP_ALIGN.CENTER, "right": PP_ALIGN.RIGHT,
    "justify": PP_ALIGN.JUSTIFY, "start": PP_ALIGN.LEFT, "end": PP_ALIGN.RIGHT,
}
ANCHOR_MAP = {"top": MSO_ANCHOR.TOP, "center": MSO_ANCHOR.MIDDLE, "bottom": MSO_ANCHOR.BOTTOM}


# --------------------------------------------------------------------------
# XML 辅助
# --------------------------------------------------------------------------

def set_alpha(color_el, alpha: float) -> None:
    """在 srgbClr 上写 alpha（0~1）。"""
    if color_el is None or alpha >= 0.999:
        return
    srgb = color_el.find(qn("a:srgbClr"))
    if srgb is None:
        return
    for old in srgb.findall(qn("a:alpha")):
        srgb.remove(old)
    node = srgb.makeelement(qn("a:alpha"), {"val": str(int(max(0.0, min(1.0, alpha)) * 100000))})
    srgb.append(node)


def apply_alpha_to_fill(shape, alpha: float) -> None:
    if alpha >= 0.999:
        return
    try:
        solid = shape.fill._xPr.find(qn("a:solidFill"))
        if solid is not None:
            set_alpha(solid, alpha)
    except Exception:
        pass


def add_outer_shadow(shape, blur: float = 12.0, distance: float = 4.0, direction: float = 90.0, alpha: float = 0.25) -> None:
    """写入基础外阴影（outerShdw）。"""
    try:
        sp_pr = shape._element.spPr
        for old in sp_pr.findall(qn("a:effectLst")):
            sp_pr.remove(old)
        effect = sp_pr.makeelement(qn("a:effectLst"), {})
        shdw = effect.makeelement(
            qn("a:outerShdw"),
            {
                "blurRad": str(px_to_emu(blur)),
                "dist": str(px_to_emu(distance)),
                "dir": str(int(direction * 60000)),
                "rotWithShape": "0",
            },
        )
        clr = shdw.makeelement(qn("a:srgbClr"), {"val": "000000"})
        a = clr.makeelement(qn("a:alpha"), {"val": str(int(alpha * 100000))})
        clr.append(a)
        shdw.append(clr)
        effect.append(shdw)
        sp_pr.append(effect)
    except Exception as exc:  # pragma: no cover
        log.debug("shadow failed: %s", exc)


def props_of(box: "Box", key: str) -> Dict[str, Any]:
    """读取组件属性中的对象（如 defaultTextStyle / defaultCellStyle）。"""
    value = box.el.props.get(key)
    return value if isinstance(value, dict) else {}


def set_hyperlink(element, url: str) -> None:
    if not url:
        return
    try:
        element.click_action.hyperlink.address = url
    except Exception:
        try:
            element.hyperlink.address = url  # type: ignore[attr-defined]
        except Exception as exc:  # pragma: no cover
            log.debug("hyperlink failed: %s", exc)


# --------------------------------------------------------------------------
# 渲染器
# --------------------------------------------------------------------------

class RenderError(ValueError):
    """A required component could not be rendered faithfully."""


class Renderer:
    def __init__(self, settings: Optional[Settings] = None, base_dir: Optional[Path] = None) -> None:
        self.settings = settings or Settings(workspace=Path(base_dir or Path.cwd()))
        self.base_dir = Path(base_dir) if base_dir else Path.cwd()
        self._assets: Dict[str, bytes] = {}

    # -- deck -------------------------------------------------------------
    def new_presentation(self):
        prs = Presentation()
        prs.slide_width = Emu(px_to_emu(CANVAS_W))
        prs.slide_height = Emu(px_to_emu(CANVAS_H))
        return prs

    def render_slide_into(self, slide, slide_el: Element) -> Dict[str, Any]:
        """把一页 DSL 渲染进**已存在**的 pptx 幻灯片（先清空原有形状）。

        供 `slidep upsert-dsl` 使用：只有渲染进已有页，才能保留同文件的其他页。
        """
        for shape in list(slide.shapes):
            shape._element.getparent().remove(shape._element)
        result = layout_slide(slide_el, self.base_dir)
        bg = self._slide_background(slide, slide_el)
        self._render_container(slide, result.slide)
        self._slide_notes(slide, slide_el)
        return {
            "content_height": round(result.content_height, 1),
            "available_height": round(result.available_height, 1),
            "overflow": round(result.overflow, 1),
            "warnings": result.warnings,
            "background": bg,
        }

    def render_deck(
        self,
        slides: Sequence[Element],
        out_path: Path,
        base_dir: Optional[Path] = None,
    ) -> Dict[str, Any]:
        if base_dir:
            self.base_dir = Path(base_dir)
        prs = self.new_presentation()
        blank = prs.slide_layouts[6]
        report: List[Dict[str, Any]] = []
        for index, slide_el in enumerate(slides):
            slide = prs.slides.add_slide(blank)
            item = self.render_slide_into(slide, slide_el)
            report.append({"index": index, **item})
        out_path.parent.mkdir(parents=True, exist_ok=True)
        prs.save(str(out_path))
        return {
            "path": str(out_path),
            "slides": len(slides),
            "size_bytes": out_path.stat().st_size if out_path.exists() else 0,
            "unit": "px",
            "canvas": {"width": CANVAS_W, "height": CANVAS_H},
            "per_slide": report,
        }

    # -- 背景 / 备注 -------------------------------------------------------
    def _slide_background(self, slide, el: Element) -> str:
        style = el.style
        bg = style.get("background") or style.get("backgroundColor")
        grad = parse_gradient(bg)
        if grad:
            self._fill_gradient(slide.background.fill, grad)
            return "gradient"
        color = parse_color(bg)
        if color.is_visible():
            slide.background.fill.solid()
            slide.background.fill.fore_color.rgb = RGBColor.from_string(color.rgb_hex)
            return f"solid:{color.rgb_hex}"
        return "none"

    def _slide_notes(self, slide, el: Element) -> None:
        notes = el.props.get("notes")
        if not notes:
            return
        try:
            slide.notes_slide.notes_text_frame.text = str(notes)
        except Exception as exc:  # pragma: no cover
            log.debug("notes failed: %s", exc)

    # -- 容器 -------------------------------------------------------------
    def _render_container(self, slide, box: Box) -> None:
        self._draw_box_background(slide, box)
        for child in box.children:
            self._render_box(slide, child)

    def _draw_box_background(self, slide, box: Box) -> None:
        style = box.style
        bg = style.get("background") or style.get("backgroundColor")
        border = parse_border(style.get("border"))
        radius = resolve_length(style.get("borderRadius"), min(box.w, box.h) / 2.0, 0.0) or 0.0
        shadow = style.get("boxShadow")
        needs_shape = bool(
            (bg and parse_color(bg).is_visible())
            or parse_gradient(bg)
            or (border and border.visible())
            or shadow
        )
        if not needs_shape and not self._has_overflow_clip(style):
            return
        if not needs_shape:
            return
        is_round = radius > 0.5
        shape_type = MSO_SHAPE.ROUNDED_RECTANGLE if is_round else MSO_SHAPE.RECTANGLE
        shape = slide.shapes.add_shape(
            shape_type, Emu(px_to_emu(box.x)), Emu(px_to_emu(box.y)),
            Emu(px_to_emu(max(1.0, box.w))), Emu(px_to_emu(max(1.0, box.h))),
        )
        if is_round:
            adj = min(0.5, max(0.0, radius / max(1.0, min(box.w, box.h))))
            try:
                shape.adjustments[0] = adj
            except Exception:
                pass
        grad = parse_gradient(bg)
        color = parse_color(bg)
        if grad:
            self._fill_gradient(shape.fill, grad)
        elif color.is_visible():
            shape.fill.solid()
            shape.fill.fore_color.rgb = RGBColor.from_string(color.rgb_hex)
            apply_alpha_to_fill(shape, color.alpha)
        else:
            shape.fill.background()
        if border and border.visible():
            shape.line.color.rgb = RGBColor.from_string(border.color.rgb_hex)
            shape.line.width = Pt(max(0.25, border.width * 0.75))
        else:
            shape.line.fill.background()
        opacity = style.get("opacity")
        if isinstance(opacity, (int, float)) and float(opacity) < 0.999:
            apply_alpha_to_fill(shape, float(opacity))
        if shadow:
            add_outer_shadow(shape)
        shape.shadow.inherit = False
        set_hyperlink(shape, style.get("href") or box.el.props.get("href"))
        box.meta["shape"] = shape

    def _has_overflow_clip(self, style: Dict[str, Any]) -> bool:
        return str(style.get("overflow", "")).lower() in ("hidden", "clip")

    # -- 叶子 -------------------------------------------------------------
    def _render_box(self, slide, box: Box) -> None:
        kind = box.kind
        if kind == "text":
            self._render_text(slide, box)
        elif kind == "image":
            self._render_image(slide, box)
        elif kind == "table":
            self._render_table(slide, box)
        elif kind == "chart":
            self._render_chart(slide, box)
        elif kind == "icon":
            self._render_icon(slide, box)
        elif kind == "qrcode":
            self._render_qrcode(slide, box)
        elif kind == "code":
            self._render_code(slide, box)
        elif kind == "svg":
            self._render_svg(slide, box)
        elif kind == "diagram":
            self._render_diagram(slide, box)
        else:
            self._render_container(slide, box)

    # -- 文本 -------------------------------------------------------------
    def _render_text(self, slide, box: Box) -> None:
        style = box.style
        runs = runs_of(box.el)
        if not runs:
            return
        padding = parse_padding(style.get("padding"))
        x = box.x + padding[3]
        y = box.y + padding[0]
        w = max(4.0, box.w - padding[1] - padding[3])
        h = max(4.0, box.h - padding[0] - padding[2])
        tb = slide.shapes.add_textbox(Emu(px_to_emu(x)), Emu(px_to_emu(y)), Emu(px_to_emu(w)), Emu(px_to_emu(h)))
        tf = tb.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        tf.vertical_anchor = ANCHOR_MAP.get(str(style.get("alignItems", "top")).lower(), MSO_ANCHOR.TOP)
        base_size = float(style.get("fontSize") or 16)
        paragraph = tf.paragraphs[0]
        paragraph.alignment = ALIGN_MAP.get(str(style.get("textAlign", "left")).lower(), PP_ALIGN.LEFT)
        first = True
        for text, run_style in runs:
            if text == "\n":
                paragraph = tf.add_paragraph()
                paragraph.alignment = ALIGN_MAP.get(str(style.get("textAlign", "left")).lower(), PP_ALIGN.LEFT)
                first = True
                continue
            if not text:
                continue
            target = paragraph
            if not first and style.get("whiteSpace") == "nowrap":
                target = paragraph
            run = target.add_run()
            run.text = text
            self._style_run(run, style, run_style, base_size)
            first = False
        href = box.el.props.get("href") or style.get("href")
        if href:
            for para in tf.paragraphs:
                for run in para.runs:
                    set_hyperlink(run, str(href))
        tb.shadow.inherit = False
        box.meta["shape"] = tb

    def _style_run(self, run, base_style: Dict[str, Any], override: Dict[str, Any], base_size: float) -> None:
        merged = dict(base_style)
        merged.update(override or {})
        size = merged.get("fontSize") or base_size
        run.font.size = Pt(float(resolve_length(size, base_size, base_size) or base_size) * 0.75)
        weight = str(merged.get("fontWeight", "")).lower()
        run.font.bold = weight in ("bold", "bolder", "600", "700", "800", "900") or merged.get("bold") is True
        style_name = str(merged.get("fontStyle", "")).lower()
        run.font.italic = style_name == "italic" or merged.get("italic") is True
        decoration = str(merged.get("textDecoration", "")).lower()
        run.font.underline = "underline" in decoration
        color = parse_color(merged.get("color"))
        if color.is_visible():
            run.font.color.rgb = RGBColor.from_string(color.rgb_hex)
        family = merged.get("fontFamily") or merged.get("font")
        if family:
            run.font.name = str(family)
            self._set_east_asian_font(run, str(family))
        letter = merged.get("letterSpacing")
        if letter is not None:
            spacing = resolve_length(letter, 0.0, 0.0) or 0.0
            if abs(spacing) > 0.01:
                run.font._rPr.set("spc", str(int(spacing * 75)))

    @staticmethod
    def _set_east_asian_font(run, family: str) -> None:
        try:
            rPr = run.font._rPr
            for tag in ("a:ea", "a:cs"):
                el = rPr.find(qn(tag))
                if el is None:
                    el = rPr.makeelement(qn(tag), {})
                    rPr.append(el)
                el.set("typeface", family)
        except Exception:
            pass

    # -- 图片 -------------------------------------------------------------
    def _render_image(self, slide, box: Box) -> None:
        src = box.el.props.get("src")
        if not src:
            self._placeholder(slide, box, "[image: missing src]")
            return
        data = self._load_asset(str(src))
        if data is None:
            self._placeholder(slide, box, f"[image not found: {src}]")
            return
        try:
            stream = io.BytesIO(data)
            pic = slide.shapes.add_picture(
                stream, Emu(px_to_emu(box.x)), Emu(px_to_emu(box.y)),
                Emu(px_to_emu(max(1.0, box.w))), Emu(px_to_emu(max(1.0, box.h))),
            )
        except Exception as exc:
            self._placeholder(slide, box, f"[image decode failed: {exc}]")
            return
        radius = resolve_length(box.style.get("borderRadius"), 0.0, 0.0) or 0.0
        if radius > 0.5:
            self._round_picture(pic, radius, box)
        set_hyperlink(pic, str(box.el.props.get("href") or box.style.get("href") or ""))
        box.meta["shape"] = pic

    def _round_picture(self, pic, radius: float, box: Box) -> None:
        try:
            sp_pr = pic._element.spPr
            prst = sp_pr.makeelement(qn("a:prstGeom"), {"prst": "roundRect"})
            av = prst.makeelement(qn("a:avLst"), {})
            gd = av.makeelement(
                qn("a:gd"),
                {"name": "adj", "fmla": f"val {int(min(50000, radius / max(1.0, min(box.w, box.h)) * 100000))}"},
            )
            av.append(gd)
            prst.append(av)
            old = sp_pr.find(qn("a:prstGeom"))
            if old is not None:
                sp_pr.remove(old)
            sp_pr.insert(0, prst)
        except Exception:
            pass

    def _load_asset(self, src: str) -> bytes:
        if src.startswith(("http://", "https://")):
            raise RenderError("Remote images must first be saved as local project assets.")
        root = self.base_dir.resolve()
        path = (root / src).resolve()
        try:
            path.relative_to(root)
        except ValueError as exc:
            raise RenderError("Image source escapes the project directory.") from exc
        if not path.is_file():
            raise RenderError(f"Image not found: {src}")
        before = path.stat()
        if before.st_size > MAX_ASSET_BYTES:
            raise RenderError("Image exceeds the 16 MiB asset budget.")
        key = f"{path}:{before.st_mtime_ns}:{before.st_size}"
        if key in self._assets:
            return self._assets[key]
        with path.open("rb") as stream:
            data = stream.read(MAX_ASSET_BYTES + 1)
        after = path.stat()
        if len(data) > MAX_ASSET_BYTES or (before.st_size, before.st_mtime_ns) != (after.st_size, after.st_mtime_ns):
            raise RenderError("Image changed while being read or exceeds the asset budget.")
        with Image.open(io.BytesIO(data)) as decoded:
            if decoded.width * decoded.height > 40_000_000:
                raise RenderError("Image exceeds the 40 million pixel budget.")
        self._assets[key] = data
        return data

    # -- 表格 -------------------------------------------------------------
    def _render_table(self, slide, box: Box) -> None:
        cells = box.el.props.get("cells") or []
        if not isinstance(cells, list) or not cells:
            self._placeholder(slide, box, "[table: empty cells]")
            return
        rows = len(cells)
        cols = max(len(r) for r in cells if isinstance(r, list)) if rows else 0
        if cols == 0:
            self._placeholder(slide, box, "[table: empty cells]")
            return
        shape = slide.shapes.add_table(
            rows, cols,
            Emu(px_to_emu(box.x)), Emu(px_to_emu(box.y)),
            Emu(px_to_emu(max(1.0, box.w))), Emu(px_to_emu(max(1.0, box.h))),
        )
        table = shape.table
        defaults = props_of(box, "defaultTextStyle")
        default_cell_style = props_of(box, "defaultCellStyle")
        default_bg = self._table_bg(default_cell_style)
        base_size = float(defaults.get("fontSize") or 16)
        default_align = ALIGN_MAP.get(str(defaults.get("textAlign", "center")).lower(), PP_ALIGN.CENTER)
        for r in range(rows):
            row_cells = cells[r] if isinstance(cells[r], list) else []
            for c in range(cols):
                cell = table.cell(r, c)
                cell.margin_left = cell.margin_right = Emu(px_to_emu(6))
                cell.margin_top = cell.margin_bottom = Emu(px_to_emu(3))
                raw = row_cells[c] if c < len(row_cells) else ""
                text, text_style, cell_style = self._cell_parts(raw)
                tf = cell.text_frame
                tf.word_wrap = True
                para = tf.paragraphs[0]
                para.alignment = ALIGN_MAP.get(
                    str(text_style.get("textAlign") or defaults.get("textAlign", "center")).lower(),
                    default_align,
                )
                run = para.add_run()
                run.text = text
                size = text_style.get("fontSize") or base_size
                run.font.size = Pt(float(resolve_length(size, base_size, base_size) or base_size) * 0.75)
                weight = str(text_style.get("fontWeight", "")).lower()
                run.font.bold = bool(
                    text_style.get("bold") or weight in ("bold", "bolder", "600", "700", "800", "900")
                )
                color = parse_color(text_style.get("color") or defaults.get("color"))
                if color.is_visible():
                    run.font.color.rgb = RGBColor.from_string(color.rgb_hex)
                family = text_style.get("fontFamily") or defaults.get("fontFamily") or defaults.get("font")
                if family:
                    run.font.name = str(family)
                    self._set_east_asian_font(run, str(family))
                # 背景：单元格自身 > defaultCellStyle；都未指定则不填充
                bg = self._table_bg(cell_style)
                if bg is None:
                    bg = default_bg
                cell_color = parse_color(bg) if bg else Color(transparent=True)
                if cell_color.is_visible():
                    cell.fill.solid()
                    cell.fill.fore_color.rgb = RGBColor.from_string(cell_color.rgb_hex)
                else:
                    cell.fill.background()
        box.meta["shape"] = shape

    @staticmethod
    def _cell_parts(raw: Any) -> Tuple[str, Dict[str, Any], Dict[str, Any]]:
        """单元格支持字符串或对象格式 {text, textStyle, cellStyle}。"""
        if isinstance(raw, dict):
            text = raw.get("text")
            return ("" if text is None else str(text)), (raw.get("textStyle") or {}), (raw.get("cellStyle") or {})
        return ("" if raw is None else str(raw)), {}, {}

    @staticmethod
    def _table_bg(cell_style: Any) -> Optional[str]:
        """提取单元格背景：支持 {background:'#fff'} 与 {background:{color:'#fff'}}。"""
        if not isinstance(cell_style, dict):
            return None
        bg = cell_style.get("background") or cell_style.get("backgroundColor")
        if isinstance(bg, dict):
            bg = bg.get("color")
        return str(bg) if bg else None

    # -- 图表 -------------------------------------------------------------
    def _render_chart(self, slide, box: Box) -> None:
        props = box.el.props
        chart_type = str(props.get("chartType") or props.get("type") or "barChart")
        bar_direction = str(props.get("barDirection") or "column").lower()
        grouping = str(props.get("grouping") or "").lower()
        ctype = resolve_chart_type(chart_type, bar_direction, grouping)
        categories, series = self._chart_data(props)
        if not series:
            self._placeholder(slide, box, "[chart: no data]")
            return
        if not categories:
            categories = [str(i + 1) for i in range(max(len(v) for _, v in series))]
        width = len(categories)
        chart_data = CategoryChartData()
        chart_data.categories = categories
        for name, values in series:
            padded = (values + [0.0] * width)[:width]
            chart_data.add_series(name, padded)
        try:
            frame = slide.shapes.add_chart(
                ctype, Emu(px_to_emu(box.x)), Emu(px_to_emu(box.y)),
                Emu(px_to_emu(max(1.0, box.w))), Emu(px_to_emu(max(1.0, box.h))), chart_data,
            )
        except Exception as exc:
            self._placeholder(slide, box, f"[chart failed: {exc}]")
            return
        self._style_chart(frame.chart, props)
        box.meta["shape"] = frame

    def _chart_data(self, props: Dict[str, Any]) -> Tuple[List[str], List[Tuple[str, List[float]]]]:
        """解析两种数据写法（规范优先）：

        1) 规范写法：`data` 为二维数组，首行是表头（首列=分类名，其余列=系列名）
        2) 兼容写法：`series=[{name,data}]`（或 dict）+ 可选 `categories`
        """
        raw = props.get("data")
        categories: List[str] = []
        series: List[Tuple[str, List[float]]] = []
        if isinstance(raw, list) and raw and isinstance(raw[0], list):
            header = [str(x) for x in raw[0]]
            names = header[1:] or ["数值"]
            rows = [r for r in raw[1:] if isinstance(r, list)]
            column_count = max((len(r) - 1 for r in rows), default=len(names))
            column_count = max(column_count, len(names))
            if len(names) < column_count:
                names = names + [f"系列{i}" for i in range(len(names) + 1, column_count + 1)]
            names = names[:column_count]
            buckets: List[List[float]] = [[] for _ in names]
            for row in rows:
                categories.append(str(row[0]) if row else "")
                for index in range(len(names)):
                    value = row[index + 1] if index + 1 < len(row) else 0
                    buckets[index].append(self._to_float(value))
            series = [(names[i], buckets[i]) for i in range(len(names))]
            return categories, series

        raw_series = props.get("series")
        if isinstance(raw_series, list) and raw_series and isinstance(raw_series[0], dict):
            for item in raw_series:
                name = str(item.get("name") or f"Series {len(series) + 1}")
                values = item.get("data") or item.get("values") or []
                series.append((name, [self._to_float(v) for v in values]))
        elif isinstance(raw_series, dict):
            for name, values in raw_series.items():
                series.append((str(name), [self._to_float(v) for v in (values or [])]))
        categories = [str(c) for c in (props.get("categories") or props.get("labels") or [])]
        return categories, series

    def _style_chart(self, chart, props: Dict[str, Any]) -> None:
        """标题 / 图例 / 系列配色 / 数据标签 / 背景（对齐 component-chart 规范）。"""
        if props.get("title"):
            chart.has_title = True
            chart.chart_title.text_frame.text = str(props.get("title"))
            title_color = parse_color(props.get("titleColor"))
            if title_color.is_visible():
                for para in chart.chart_title.text_frame.paragraphs:
                    for run in para.runs:
                        run.font.color.rgb = RGBColor.from_string(title_color.rgb_hex)
        else:
            chart.has_title = False

        show_legend = props.get("showLegend")
        legend_prop = str(props.get("legend", "")).lower()
        if show_legend is None:
            show_legend = legend_prop not in ("none", "false", "off")
        chart.has_legend = bool(show_legend)
        if chart.has_legend:
            position = str(props.get("legendPosition") or legend_prop or "bottom").lower()
            chart.legend.position = {
                "top": XL_LEGEND_POSITION.TOP, "bottom": XL_LEGEND_POSITION.BOTTOM,
                "left": XL_LEGEND_POSITION.LEFT, "right": XL_LEGEND_POSITION.RIGHT,
            }.get(position, XL_LEGEND_POSITION.BOTTOM)
            chart.legend.include_in_layout = False
            legend_color = parse_color(props.get("legendColor"))
            if legend_color.is_visible():
                try:
                    chart.legend.font.color.rgb = RGBColor.from_string(legend_color.rgb_hex)
                except Exception:
                    pass

        colors = props.get("colors")
        if isinstance(colors, list):
            for index, plot_series in enumerate(chart.series):
                if index >= len(colors):
                    break
                gradient = parse_gradient(colors[index])
                color = gradient.first() if gradient else parse_color(colors[index])
                if not color.is_visible():
                    continue
                try:
                    plot_series.format.fill.solid()
                    plot_series.format.fill.fore_color.rgb = RGBColor.from_string(color.rgb_hex)
                except Exception:
                    pass

        if props.get("showDataLabels"):
            try:
                plot = chart.plots[0]
                plot.has_data_labels = True
                label_color = parse_color(props.get("dataLabelColor"))
                for label in plot.data_labels:
                    label.font.size = Pt(float(props.get("dataLabelSize") or 10))
                    if label_color.is_visible():
                        label.font.color.rgb = RGBColor.from_string(label_color.rgb_hex)
            except Exception:
                pass

        background = props.get("background")
        if background:
            gradient = parse_gradient(background)
            color = gradient.first() if gradient else parse_color(background)
            if color.is_visible():
                try:
                    chart.plots[0].format.fill.solid()
                    chart.plots[0].format.fill.fore_color.rgb = RGBColor.from_string(color.rgb_hex)
                except Exception:
                    pass

    @staticmethod
    def _to_float(value: Any) -> float:
        try:
            return float(value)
        except (TypeError, ValueError):
            return 0.0

    # -- 图标 -------------------------------------------------------------
    def _render_icon(self, slide, box: Box) -> None:
        name = str(box.el.props.get("name") or "").lower()
        style = box.style
        fill = parse_color(box.el.props.get("fill") or style.get("fill") or style.get("color"))
        if not fill.is_visible():
            self._placeholder(slide, box, f"[icon {name}: fill required]")
            return
        shape_type = ICON_SHAPES.get(name)
        if shape_type is None:
            glyph = ICON_GLYPHS.get(name, name[:1].upper() if name else "?")
            tb = slide.shapes.add_textbox(
                Emu(px_to_emu(box.x)), Emu(px_to_emu(box.y)),
                Emu(px_to_emu(max(1.0, box.w))), Emu(px_to_emu(max(1.0, box.h))),
            )
            tf = tb.text_frame
            tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
            para = tf.paragraphs[0]
            para.alignment = PP_ALIGN.CENTER
            run = para.add_run()
            run.text = glyph
            run.font.size = Pt(max(6.0, min(box.w, box.h) * 0.75))
            run.font.bold = True
            run.font.color.rgb = RGBColor.from_string(fill.rgb_hex)
            tb.shadow.inherit = False
            box.meta["shape"] = tb
            return
        shape = slide.shapes.add_shape(
            shape_type, Emu(px_to_emu(box.x)), Emu(px_to_emu(box.y)),
            Emu(px_to_emu(max(1.0, box.w))), Emu(px_to_emu(max(1.0, box.h))),
        )
        shape.fill.solid()
        shape.fill.fore_color.rgb = RGBColor.from_string(fill.rgb_hex)
        apply_alpha_to_fill(shape, fill.alpha)
        shape.line.fill.background()
        shape.shadow.inherit = False
        if name == "check-circle":
            shape.fill.solid()
            shape.fill.fore_color.rgb = RGBColor.from_string("FFFFFF")
            shape.line.color.rgb = RGBColor.from_string(fill.rgb_hex)
            shape.line.width = Pt(2.0)
        glyph = ICON_GLYPHS.get(name)
        if glyph:
            tf = shape.text_frame
            tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
            para = tf.paragraphs[0]
            para.alignment = PP_ALIGN.CENTER
            run = para.add_run()
            run.text = glyph
            run.font.bold = True
            run.font.size = Pt(max(6.0, min(box.w, box.h) * 0.6))
            run.font.color.rgb = RGBColor.from_string(fill.rgb_hex if name != "check-circle" else fill.rgb_hex)
        box.meta["shape"] = shape

    # -- 二维码 -----------------------------------------------------------
    def _render_qrcode(self, slide, box: Box) -> None:
        text = str(box.el.props.get("text") or "")
        if not text:
            self._placeholder(slide, box, "[qrcode: text required]")
            return
        size = float(box.w or box.el.props.get("width") or 160)
        dark = parse_color(box.el.props.get("darkColor") or "#000000")
        light = parse_color(box.el.props.get("lightColor") or "#FFFFFF")
        ecc = str(box.el.props.get("errorCorrectionLevel") or "M").upper()[:1]
        try:
            from qrcode.constants import (
                ERROR_CORRECT_H,
                ERROR_CORRECT_L,
                ERROR_CORRECT_M,
                ERROR_CORRECT_Q,
            )

            qr = qrcode.QRCode(
                error_correction={
                    "L": ERROR_CORRECT_L, "M": ERROR_CORRECT_M, "Q": ERROR_CORRECT_Q, "H": ERROR_CORRECT_H,
                }.get(ecc, ERROR_CORRECT_M),
                box_size=10,
                border=2,
            )
            qr.add_data(text)
            qr.make(fit=True)
            img = qr.make_image(
                fill_color=f"#{dark.rgb_hex}", back_color=f"#{light.rgb_hex}"
            )
            stream = io.BytesIO()
            img.save(stream, format="PNG")
            stream.seek(0)
            pic = slide.shapes.add_picture(
                stream, Emu(px_to_emu(box.x)), Emu(px_to_emu(box.y)),
                Emu(px_to_emu(size)), Emu(px_to_emu(size)),
            )
            box.meta["shape"] = pic
        except Exception as exc:
            self._placeholder(slide, box, f"[qrcode failed: {exc}]")

    # -- 代码块 -----------------------------------------------------------
    def _render_code(self, slide, box: Box) -> None:
        code = str(box.el.props.get("code") or "")
        font_size = float(box.el.props.get("fontSize") or 12)
        theme = str(box.el.props.get("theme") or "github-dark").lower()
        mac_header = box.el.props.get("macHeader", True)
        dark = "dark" in theme
        bg = parse_color(box.el.props.get("background") or ("#0D1117" if dark else "#F6F8FA"))
        fg = parse_color(box.el.props.get("color") or ("#E6EDF3" if dark else "#1F2328"))
        shape = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE, Emu(px_to_emu(box.x)), Emu(px_to_emu(box.y)),
            Emu(px_to_emu(max(1.0, box.w))), Emu(px_to_emu(max(1.0, box.h))),
        )
        try:
            shape.adjustments[0] = 0.04
        except Exception:
            pass
        shape.fill.solid()
        shape.fill.fore_color.rgb = RGBColor.from_string(bg.rgb_hex)
        shape.line.fill.background()
        shape.shadow.inherit = False
        header_h = 26.0 if mac_header else 0.0
        x = box.x + 16
        y = box.y + header_h + 8
        w = max(4.0, box.w - 32)
        if mac_header:
            for i, color in enumerate(("#FF5F56", "#FFBD2E", "#27C93F")):
                dot = slide.shapes.add_shape(
                    MSO_SHAPE.OVAL, Emu(px_to_emu(box.x + 14 + i * 16)),
                    Emu(px_to_emu(box.y + 9)), Emu(px_to_emu(9)), Emu(px_to_emu(9)),
                )
                dot.fill.solid()
                dot.fill.fore_color.rgb = RGBColor.from_string(color.lstrip("#"))
                dot.line.fill.background()
                dot.shadow.inherit = False
        tb = slide.shapes.add_textbox(Emu(px_to_emu(x)), Emu(px_to_emu(y)), Emu(px_to_emu(w)), Emu(px_to_emu(max(4.0, box.h - header_h - 12))))
        tf = tb.text_frame
        tf.word_wrap = False
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        lines = code.replace("\t", "    ").splitlines() or [""]
        for index, line in enumerate(lines[:80]):
            para = tf.paragraphs[0] if index == 0 else tf.add_paragraph()
            run = para.add_run()
            run.text = line if line else " "
            run.font.size = Pt(font_size * 0.75)
            run.font.name = "Consolas"
            run.font.color.rgb = RGBColor.from_string(fg.rgb_hex)
            self._set_east_asian_font(run, "Consolas")
        tb.shadow.inherit = False
        box.meta["shape"] = shape

    # -- SVG --------------------------------------------------------------
    def _render_svg(self, slide, box: Box) -> None:
        markup = self._svg_markup(box.el)
        png = rasterize_svg(markup, box.w, box.h)
        if png is None:
            self._placeholder(slide, box, f"[svg: {len(markup)} chars]")
            return
        stream = io.BytesIO(png)
        try:
            pic = slide.shapes.add_picture(
                stream, Emu(px_to_emu(box.x)), Emu(px_to_emu(box.y)),
                Emu(px_to_emu(max(1.0, box.w))), Emu(px_to_emu(max(1.0, box.h))),
            )
            box.meta["shape"] = pic
        except Exception as exc:
            self._placeholder(slide, box, f"[svg render failed: {exc}]")

    def _svg_markup(self, el: Element) -> str:
        raw = el.props.get("content") or el.props.get("markup")
        if raw:
            return str(raw)
        width = el.props.get("width") or 100
        height = el.props.get("height") or 100
        viewbox = el.props.get("viewBox") or f"0 0 {width} {height}"
        parts = [f'<svg xmlns="http://www.w3.org/2000/svg" width="{width}" height="{height}" viewBox="{viewbox}">']
        for child in el.children:
            if isinstance(child, Element):
                parts.append(node_to_svg(child))
        parts.append("</svg>")
        return "".join(parts)

    # -- Diagram ----------------------------------------------------------
    def _render_diagram(self, slide, box: Box) -> None:
        raise RenderError("Diagram is unsupported; supply a verified local image.")

    def _placeholder(self, slide, box: Box, message: str) -> None:
        raise RenderError(message)


# --------------------------------------------------------------------------
# SVG 子集栅格化（无第三方渲染器时的内建回退）
# --------------------------------------------------------------------------

_SVG_TAG_RE = re.compile(r"<([A-Za-z]+)([^>]*?)/?>", re.S)
_ATTR_RE = re.compile(r"([A-Za-z_:][-\w:.]*)\s*=\s*\"([^\"]*)\"")


def node_to_svg(el: Element) -> str:
    tag = el.tag
    attrs = " ".join(
        f'{k}="{v}"' for k, v in el.props.items() if isinstance(v, (str, int, float))
    )
    if tag == "br":
        return "<br/>"
    children = "".join(node_to_svg(c) for c in el.children if isinstance(c, Element))
    return f"<{tag} {attrs}>{children}</{tag}>"


def rasterize_svg(markup: str, width: float, height: float, scale: float = 2.0) -> Optional[bytes]:
    """用 PIL 渲染 SVG 的常用子集（rect/circle/ellipse/line/polygon/path）。"""
    if not markup:
        return None
    w = max(4, int(width * scale))
    h = max(4, int(height * scale))
    img = Image.new("RGBA", (w, h), (0, 0, 0, 0))
    draw = ImageDraw.Draw(img)

    vb = re.search(r'viewBox\s*=\s*"([^"]+)"', markup)
    vb_scale = 1.0
    if vb:
        parts = [float(x) for x in re.split(r"[\s,]+", vb.group(1).strip()) if x]
        if len(parts) == 4 and parts[2] and parts[3]:
            vb_scale = min(w / parts[2], h / parts[3])

    def num(value: Any, default: float = 0.0) -> float:
        try:
            return float(re.sub(r"[a-zA-Z%]+$", "", str(value)))
        except (TypeError, ValueError):
            return default

    def paint(value: Optional[str]) -> Optional[tuple]:
        if value is None or value.lower() in ("none", "transparent", ""):
            return None
        parsed = parse_color(value)
        if not parsed.is_visible():
            return None
        rgb = tuple(int(parsed.rgb_hex[i : i + 2], 16) for i in (0, 2, 4))
        return rgb + (int(parsed.alpha * 255),)

    for match in _SVG_TAG_RE.finditer(markup):
        tag = match.group(1).lower()
        attrs = dict(_ATTR_RE.findall(match.group(2) or ""))
        if tag in ("svg", "g", "defs", "title", "desc"):
            continue
        fill = paint(attrs.get("fill"))
        stroke = paint(attrs.get("stroke"))
        stroke_w = max(1, int(num(attrs.get("stroke-width"), 1) * vb_scale * scale))

        def pt(x: float, y: float) -> tuple:
            return (x * vb_scale * scale, y * vb_scale * scale)

        if tag == "rect":
            x, y = num(attrs.get("x")), num(attrs.get("y"))
            rw, rh = num(attrs.get("width")), num(attrs.get("height"))
            draw.rectangle([pt(x, y), pt(x + rw, y + rh)], fill=fill, outline=stroke, width=stroke_w)
        elif tag in ("circle", "ellipse"):
            cx, cy = num(attrs.get("cx")), num(attrs.get("cy"))
            rx = num(attrs.get("r")) or num(attrs.get("rx"))
            ry = num(attrs.get("r")) or num(attrs.get("ry"))
            draw.ellipse([pt(cx - rx, cy - ry), pt(cx + rx, cy + ry)], fill=fill, outline=stroke, width=stroke_w)
        elif tag == "line":
            draw.line([pt(num(attrs.get("x1")), num(attrs.get("y1"))), pt(num(attrs.get("x2")), num(attrs.get("y2")))], fill=stroke or fill, width=stroke_w)
        elif tag == "polygon":
            raw = [float(v) for v in re.split(r"[\s,]+", attrs.get("points", "")) if v]
            pairs = list(zip(raw[0::2], raw[1::2]))
            if pairs:
                draw.polygon([pt(x, y) for x, y in pairs], fill=fill, outline=stroke)
        elif tag == "polyline":
            raw = [float(v) for v in re.split(r"[\s,]+", attrs.get("points", "")) if v]
            pairs = list(zip(raw[0::2], raw[1::2]))
            if len(pairs) > 1:
                draw.line([pt(x, y) for x, y in pairs], fill=stroke or fill, width=stroke_w, joint="curve")
        elif tag == "text":
            x, y = num(attrs.get("x")), num(attrs.get("y"))
            size = max(6, int(num(attrs.get("font-size"), 12) * vb_scale * scale))
            draw.text(pt(x, y - num(attrs.get("font-size"), 12) * 0.8), match.group(2) and "" or "", fill=fill, font_size=size)
        elif tag == "path":
            d = attrs.get("d", "")
            pairs = _parse_path(d)
            if len(pairs) > 1:
                draw.line([pt(x, y) for x, y in pairs], fill=stroke or fill, width=stroke_w, joint="curve")
            elif fill and pairs:
                draw.polygon([pt(x, y) for x, y in pairs], fill=fill)
    out = io.BytesIO()
    img.save(out, format="PNG")
    return out.getvalue()


_PATH_TOKEN_RE = re.compile(r"([MmLlHhVvZzCcSsQqTtAa])|(-?\d*\.?\d+)")


def _parse_path(d: str) -> List[Tuple[float, float]]:
    """极简 path 解析：支持 M/L/H/V/Z（含相对），曲线按端点近似。"""
    points: List[Tuple[float, float]] = []
    cursor = (0.0, 0.0)
    start = (0.0, 0.0)
    nums: List[float] = []
    cmd = "M"
    for match in _PATH_TOKEN_RE.finditer(d):
        if match.group(1):
            cmd = match.group(1)
            if cmd.upper() == "Z":
                cursor = start
                points.append(cursor)
            nums = []
            continue
        nums.append(float(match.group(2)))
        take = {"M": 2, "L": 2, "H": 1, "V": 1, "C": 6, "S": 4, "Q": 4, "T": 2, "A": 7}.get(cmd.upper(), 2)
        if len(nums) >= take:
            if cmd.upper() == "M":
                cursor = (nums[0], nums[1]) if cmd == "M" else (cursor[0] + nums[0], cursor[1] + nums[1])
                start = cursor
                points.append(cursor)
                cmd = "L" if cmd == "M" else "l"
            elif cmd.upper() == "L":
                cursor = (nums[0], nums[1]) if cmd == "L" else (cursor[0] + nums[0], cursor[1] + nums[1])
                points.append(cursor)
            elif cmd.upper() == "H":
                cursor = (nums[0], cursor[1]) if cmd == "H" else (cursor[0] + nums[0], cursor[1])
                points.append(cursor)
            elif cmd.upper() == "V":
                cursor = (cursor[0], nums[0]) if cmd == "V" else (cursor[0], cursor[1] + nums[0])
                points.append(cursor)
            elif cmd.upper() in ("C", "S", "Q", "T"):
                dx, dy = nums[-2], nums[-1]
                cursor = (dx, dy) if cmd.isupper() else (cursor[0] + dx, cursor[1] + dy)
                points.append(cursor)
            elif cmd.upper() == "A":
                dx, dy = nums[-2], nums[-1]
                cursor = (dx, dy) if cmd.isupper() else (cursor[0] + dx, cursor[1] + dy)
                points.append(cursor)
            nums = []
    return points


# --------------------------------------------------------------------------
# 对外 API
# --------------------------------------------------------------------------
