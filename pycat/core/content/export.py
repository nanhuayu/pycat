"""Shared Markdown document conversion. No GUI, network access or application state."""
from __future__ import annotations

import base64
import io
from html import escape
from pathlib import Path
from urllib.parse import unquote, urlsplit

from pycat.core.persistence import atomic_write_bytes

# Format keys, suffixes and display names are shared by adapters.
DOCUMENT_FORMATS = {"markdown": (".md", "Markdown"), "docx": (".docx", "Word"), "html": (".html", "HTML")}
CONVERSATION_FORMATS = {**DOCUMENT_FORMATS, "json": (".json", "JSON")}


def document_format(destination: str | Path, format: str | None = None) -> str:
    value = str(format or Path(destination).suffix.lstrip(".")).lower()
    value = {"md": "markdown", "htm": "html"}.get(value, value)
    if value not in CONVERSATION_FORMATS:
        raise ValueError(f"Unsupported export format: {value}")
    return value


def export_document(text: str, destination: str | Path, *, format: str | None = None,
                    title: str = "", source_dir: str | Path | None = None) -> Path:
    """Atomically export Markdown to MD, DOCX or self-contained HTML.

    Images are embedded only from within an explicitly supplied source directory.
    Remote images and paths outside that directory remain descriptive text.
    """
    target = Path(destination).expanduser()
    payload = render_document(text, document_format(target, format), title=title, source_dir=source_dir)
    atomic_write_bytes(target, payload)
    return target


def render_document(text: str, format: str, *, title: str = "", source_dir: str | Path | None = None) -> bytes:
    """Prepare bytes without writing, for adapters with a separate commit step."""
    if format == "markdown":
        return text.encode("utf-8")
    if format not in {"html", "docx"}:
        raise ValueError(f"Unsupported document format: {format}")
    # Conversion libraries belong to this operation, not SDK or GUI startup.
    import markdown
    from lxml import html

    tree = html.fragment_fromstring(markdown.markdown(text, extensions=["tables", "fenced_code", "sane_lists"]),
                                    create_parent="div")
    allowed = {"h1", "h2", "h3", "h4", "h5", "h6", "p", "ul", "ol", "li", "strong", "b", "em", "i",
               "del", "s", "code", "pre", "blockquote", "table", "thead", "tbody", "tr", "th", "td",
               "a", "img", "br", "hr"}
    for element in list(tree.iterdescendants()):
        if element.tag in {"script", "style", "iframe", "object", "embed"} or not isinstance(element.tag, str):
            element.drop_tree()
        elif element.tag not in allowed:
            element.drop_tag()
        else:
            attrs = dict(element.attrib)
            element.attrib.clear()
            if element.tag == "a" and urlsplit(attrs.get("href", "")).scheme in {"http", "https", "mailto"}:
                element.set("href", attrs["href"])
            elif element.tag == "img":
                element.set("alt", attrs.get("alt", "image"))
                data = _local_image(attrs.get("src", ""), source_dir)
                if data:
                    element.set("src", "data:image/png;base64," + base64.b64encode(data).decode("ascii"))
                else:
                    element.tag = "span"
                    element.text = f"[{element.get('alt')}]"
                    element.attrib.clear()
    if format == "docx":
        return _docx(tree, title)
    body = html.tostring(tree, encoding="unicode")
    return ("<!doctype html><html><head><meta charset=\"utf-8\">"
            f"<title>{escape(title or 'PyCat document')}</title><style>"
            "body{max-width:52rem;margin:2rem auto;padding:0 1rem;font:16px/1.7 system-ui,sans-serif;color:#222}"
            "h1,h2,h3{line-height:1.3}pre,blockquote{padding:1rem;background:#f5f5f7}"
            "pre{white-space:pre-wrap}table{border-collapse:collapse;width:100%}"
            "td,th{border:1px solid #ddd;padding:.4rem .6rem;text-align:left}"
            "img{max-width:100%;height:auto}body{overflow-wrap:anywhere}"
            f"</style></head><body>{body}</body></html>").encode("utf-8")


def _local_image(source: str, directory: str | Path | None) -> bytes | None:
    if directory is None or not source or urlsplit(source).scheme or source.startswith(("//", "\\\\")):
        return None
    root = Path(directory).resolve()
    path = (root / unquote(source)).resolve()
    if not path.is_relative_to(root) or not path.is_file():
        return None
    from PIL import Image

    try:
        with Image.open(path) as image:
            output = io.BytesIO()
            image.convert("RGBA").save(output, "PNG")
            return output.getvalue()
    except (OSError, ValueError):
        return None


def _docx(tree, title: str) -> bytes:
    from docx import Document
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.opc.constants import RELATIONSHIP_TYPE
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Mm, Pt, RGBColor

    document = Document()
    document.core_properties.title = title
    section = document.sections[0]
    section.page_width, section.page_height = Mm(210), Mm(297)
    section.top_margin = section.bottom_margin = Mm(20)
    section.left_margin = section.right_margin = Mm(22)
    normal = document.styles["Normal"]
    normal.font.name, normal.font.size = "Arial", Pt(10.5)
    normal.element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.paragraph_format.space_after = Pt(7)
    normal.paragraph_format.line_spacing = 1.15
    for name in ("Title", "Heading 1", "Heading 2", "Heading 3", "Heading 4", "Heading 5", "Heading 6"):
        document.styles[name].font.color.rgb = RGBColor.from_string("333333")
    code_style = document.styles.add_style("PyCat Code", 1)
    code_style.base_style = normal
    code_style.font.name, code_style.font.size = "Consolas", Pt(9)
    code_style.paragraph_format.space_before = Pt(5)

    def inline(element, paragraph, bold=False, italic=False, code=False, strike=False):
        def add(text):
            if not text:
                return
            run = paragraph.add_run(text)
            run.bold, run.italic, run.font.strike = bold, italic, strike
            if code:
                run.font.name, run.font.size = "Consolas", Pt(9)
        add(element.text)
        for child in element:
            tag = child.tag
            if tag in {"ul", "ol", "p"}:
                continue
            if tag == "br":
                paragraph.add_run().add_break()
            elif tag == "img":
                raw = base64.b64decode(child.get("src").split(",", 1)[1])
                from PIL import Image
                with Image.open(io.BytesIO(raw)) as picture:
                    width = min(166, 210 * picture.width / picture.height, picture.width * 25.4 / 96)
                paragraph.add_run().add_picture(io.BytesIO(raw), width=Mm(width))
            elif tag == "a" and child.get("href"):
                link = OxmlElement("w:hyperlink")
                link.set(qn("r:id"), paragraph.part.relate_to(child.get("href"), RELATIONSHIP_TYPE.HYPERLINK,
                                                            is_external=True))
                run = paragraph.add_run(child.text_content())
                run.font.color.rgb = RGBColor.from_string("285DA8")
                run.underline = True
                link.append(run._r)
                paragraph._p.append(link)
            else:
                inline(child, paragraph, bold or tag in {"strong", "b"}, italic or tag in {"em", "i"},
                       code or tag == "code", strike or tag in {"del", "s"})
            add(child.tail)

    def numbering():
        # Restart each ordered list, reusing Word's built-in decimal definition.
        root = document.part.numbering_part.element
        abstract = root.xpath('w:abstractNum[w:lvl/w:pStyle[@w:val="ListNumber"]]')[0]
        num = root.add_num(int(abstract.get(qn("w:abstractNumId"))))
        num.add_lvlOverride(ilvl=0).add_startOverride(1)
        return num.numId

    def block(element, container=document, depth=0):
        tag = element.tag
        if tag in {"ul", "ol"}:
            num_id = numbering() if tag == "ol" else None
            for item in element:
                paragraph = container.add_paragraph(style="List Number" if num_id else "List Bullet")
                paragraph.paragraph_format.left_indent = Mm(6 + depth * 6)
                if num_id:
                    props = paragraph._p.get_or_add_pPr().get_or_add_numPr()
                    props.get_or_add_numId().val = num_id
                    props.get_or_add_ilvl().val = 0
                inline(item, paragraph)
                for child in item:
                    if child.tag == "p":
                        if paragraph.text:
                            paragraph = container.add_paragraph()
                            paragraph.paragraph_format.left_indent = Mm(6 + depth * 6)
                        inline(child, paragraph)
                    elif child.tag in {"ul", "ol"}:
                        block(child, container, depth + 1)
            return
        if tag == "table":
            rows = element.xpath(".//tr")
            if not rows:
                return
            table = container.add_table(rows=0, cols=max(len(row) for row in rows))
            table.style, table.alignment = "Table Grid", WD_TABLE_ALIGNMENT.CENTER
            for row in rows:
                cells = table.add_row().cells
                for index, cell in enumerate(row):
                    inline(cell, cells[index].paragraphs[0], bold=cell.tag == "th")
            container.add_paragraph()
            return
        if tag == "blockquote":
            for child in element:
                paragraph = container.add_paragraph(style="Quote")
                inline(child, paragraph)
            return
        if tag == "pre":
            container.add_paragraph(element.text_content().rstrip("\n"), style="PyCat Code")
            return
        if tag == "hr":
            container.add_paragraph("—" * 12)
            return
        style = f"Heading {tag[1]}" if tag in {"h1", "h2", "h3", "h4", "h5", "h6"} else None
        inline(element, container.add_paragraph(style=style))

    if tree.text and tree.text.strip():
        document.add_paragraph(tree.text)
    for element in tree:
        block(element)
        if element.tail and element.tail.strip():
            document.add_paragraph(element.tail)
    output = io.BytesIO()
    document.save(output)
    return output.getvalue()
